"""Import observed PSA 2020 QC household counts without modifying the workbook.

The visible ``by Bgy_1`` sheet contains the whole NCR. Only the Quezon City
block is read. The hidden duplicate sheet and population projections are not
inputs. Run from any directory with ``python scripts/import_psa_households.py``.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
import sys
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from backend.services.cpdd_data_service import (
    EXPECTED_BARANGAY_COUNTS,
    PSA_HOUSEHOLD_TOTALS,
    validate_psa_household_reference,
)

DEFAULT_SOURCE = PROJECT_ROOT / (
    "data/raw/CPDD/Quezon City Educational Research Files/"
    "Demographic Data (Population)/PSA-2020-HH-BY-barangay-QC.xlsx"
)
DEFAULT_REFERENCE = PROJECT_ROOT / "data/authoritative/cpdd/district_profiles_2020.json"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/authoritative/cpdd/barangay_households_2020.json"
SOURCE_SHEET = "by Bgy_1"
FIELDS = tuple(PSA_HOUSEHOLD_TOTALS)


def source_path(path: Path) -> str:
    """Keep repository sources portable while allowing explicitly supplied paths."""
    resolved = path.resolve()
    try:
        return resolved.relative_to(PROJECT_ROOT).as_posix()
    except ValueError:
        return resolved.as_posix()


def name_key(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    text = re.sub(r"\([^)]*\)", "", text).casefold()
    key = re.sub(r"[^a-z0-9]", "", text)
    return {
        "aurora": "donaaurora",
        "santodomingo": "stodomingo",
        "balongbato": "balonbato",
        "santonino": "santonio",
    }.get(key, key)


def count_value(value: Any, location: str) -> int:
    if (
        isinstance(value, bool)
        or not isinstance(value, (int, float))
        or not math.isfinite(value)
        or value < 0
        or int(value) != value
    ):
        raise ValueError(f"{location}: expected a non-negative integer count, got {value!r}")
    return int(value)


def read_qc_block(path: Path) -> tuple[list[dict[str, Any]], dict[str, int], int]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Required source sheet {SOURCE_SHEET!r} is missing")
        sheet = workbook[SOURCE_SHEET]
        if sheet.sheet_state != "visible":
            raise ValueError("The authoritative source sheet must be visible")
        values = list(sheet.iter_rows(min_col=1, max_col=4, values_only=True))
        title = " ".join(str(v) for row in values[:3] for v in row if v is not None)
        if not re.search(r"(?:01 May 2020|May 1, 2020)", title, flags=re.I):
            raise ValueError("Source must identify the actual May 1, 2020 census date")
        if "project" in title.casefold():
            raise ValueError("Population projections are not an observed census source")
        expected_headers = ("total population", "household population", "number of households")
        actual_headers = tuple(
            " ".join(str(values[r][c] or "").strip() for r in (4, 5)).casefold()
            for c in (1, 2, 3)
        )
        if actual_headers != expected_headers:
            raise ValueError(f"Unexpected population/household headers: {actual_headers!r}")
        city_rows = [
            i for i, row in enumerate(values)
            if isinstance(row[0], str) and row[0].strip().casefold() == "quezon city"
        ]
        if len(city_rows) != 1:
            raise ValueError("Expected exactly one Quezon City control row")
        start = city_rows[0]
        totals = {
            field: count_value(values[start][i + 1], f"{SOURCE_SHEET}!{chr(66+i)}{start+1}")
            for i, field in enumerate(FIELDS)
        }
        records = []
        for index in range(start + 1, len(values)):
            row = values[index]
            if all(v is None for v in row):
                break
            if isinstance(row[0], str) and row[0].strip().upper().startswith("CITY OF "):
                break
            if not isinstance(row[0], str) or not row[0].strip():
                raise ValueError(f"Missing barangay name at {SOURCE_SHEET}!A{index+1}")
            records.append({
                "source_barangay_name": row[0].strip(),
                "source_cells": f"A{index+1}:D{index+1}",
                **{
                    field: count_value(row[i + 1], f"{SOURCE_SHEET}!{chr(66+i)}{index+1}")
                    for i, field in enumerate(FIELDS)
                },
            })
        return records, totals, start + 1
    finally:
        workbook.close()


def build_reference(source: Path, cpdd_path: Path) -> dict[str, Any]:
    cpdd = json.loads(cpdd_path.read_text(encoding="utf-8"))
    if cpdd.get("reference_year") != 2020 or cpdd.get("is_authoritative_reference") is not True:
        raise ValueError("Canonical mapping must be an authoritative 2020 CPDD reference")
    canonical = {}
    districts = cpdd.get("districts", {})
    if set(districts) != set(EXPECTED_BARANGAY_COUNTS):
        raise ValueError("Canonical mapping must contain Districts I-VI")
    for district_id, expected_count in EXPECTED_BARANGAY_COUNTS.items():
        rows = districts[district_id].get("barangays", [])
        if len(rows) != expected_count:
            raise ValueError(f"Canonical District {district_id} must have {expected_count} barangays")
        for row in rows:
            name = row.get("barangay_name")
            if not isinstance(name, str) or not name.strip() or name_key(name) in canonical:
                raise ValueError("Canonical mapping has missing or duplicate barangay names")
            count_value(row.get("population_2020"), f"CPDD {name}")
            canonical[name_key(name)] = {**row, "district_id": district_id}

    records, source_totals, city_row = read_qc_block(source)
    if source_totals != PSA_HOUSEHOLD_TOTALS:
        raise ValueError(f"Unexpected PSA city control totals: {source_totals}")
    matched = {}
    for row in records:
        key = name_key(row["source_barangay_name"])
        if key not in canonical:
            raise ValueError(f"Unmapped PSA barangay: {row['source_barangay_name']!r}")
        if key in matched:
            raise ValueError(f"Duplicate PSA barangay: {row['source_barangay_name']!r}")
        match = canonical[key]
        if row["population_2020"] != match["population_2020"]:
            raise ValueError(f"PSA/CPDD population mismatch for {match['barangay_name']}")
        if row["household_population_2020"] > row["population_2020"]:
            raise ValueError(f"Household population exceeds total population for {match['barangay_name']}")
        households = row["households_2020"]
        household_population = row["household_population_2020"]
        if households == 0 and household_population != 0:
            raise ValueError(f"Positive household population with no households for {match['barangay_name']}")
        if households > household_population:
            raise ValueError(f"Household count exceeds household population for {match['barangay_name']}")
        matched[key] = {
            "barangay_name": match["barangay_name"],
            "district_id": match["district_id"],
            **row,
            "average_household_size": round(household_population / households, 6) if households else None,
        }
        # The supplied workbook has no pcode column. Preserve canonical codes
        # if a later mapping supplies them; never invent codes from row numbers.
        for field in ("pcode", "ADM4_PCODE", "psgc_code"):
            if match.get(field) is not None:
                matched[key][field] = match[field]
    missing = set(canonical) - set(matched)
    if missing:
        raise ValueError(f"Missing PSA barangays: {', '.join(canonical[k]['barangay_name'] for k in sorted(missing))}")

    ordered = [matched[key] for key in canonical]
    summaries = {}
    for district_id in EXPECTED_BARANGAY_COUNTS:
        members = [row for row in ordered if row["district_id"] == district_id]
        totals = {field: sum(row[field] for row in members) for field in FIELDS}
        summaries[district_id] = {
            "barangay_count": len(members), **totals,
            "average_household_size": round(totals["household_population_2020"] / totals["households_2020"], 6) if totals["households_2020"] else None,
        }
    result = {
        "dataset_id": "qc_psa_barangay_households_2020",
        "title": "Quezon City Barangay Population and Households, 2020 Census",
        "is_sample": False,
        "is_authoritative_reference": True,
        "reference_year": 2020,
        "reference_date": "2020-05-01",
        "observation_type": "census_actual",
        "geographic_scope": "Quezon City",
        "publisher": "Philippine Statistics Authority (PSA)",
        "provided_by": "Quezon City Planning and Development Department (CPDD)",
        "source_document": {
            "path": source_path(source), "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "sheet": SOURCE_SHEET, "city_control_cells": f"A{city_row}:D{city_row}",
            "barangay_cells": f"A{city_row+1}:D{city_row+len(records)}",
        },
        "geographic_mapping": {
            "dataset_id": cpdd["dataset_id"], "path": source_path(cpdd_path),
            "sha256": hashlib.sha256(cpdd_path.read_bytes()).hexdigest(), "reference_year": 2020,
        },
        "field_definitions": {
            "population_2020": {"unit": "persons", "definition": "Total population, including institutional population."},
            "household_population_2020": {"unit": "persons", "definition": "Household population, excluding institutional population."},
            "households_2020": {"unit": "households", "definition": "Number of households."},
            "average_household_size": {
                "unit": "persons per household", "is_derived": True,
                "formula": "household_population_2020 / households_2020", "decimal_places": 6,
            },
        },
        "notes": [
            "Observed census reference data. This import does not train a prediction model.",
            "Only the visible by Bgy_1 Quezon City block is imported; duplicate sheets and other cities are excluded.",
            "Population projections, land area and density are not combined with these census observations.",
            "Average household size uses household population, not total population; aggregate ratios use aggregate counts.",
            "The source workbook has no pcode column. Any retained geographic codes come from the canonical mapping.",
        ],
        "city_summary": {
            "district_count": 6, "barangay_count": len(ordered), **source_totals,
            "average_household_size": round(source_totals["household_population_2020"] / source_totals["households_2020"], 6),
        },
        "district_summaries": summaries,
        "barangays": ordered,
    }
    validate_psa_household_reference(result, cpdd)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--cpdd-reference", type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if args.output.resolve() in {args.source.resolve(), args.cpdd_reference.resolve()}:
        parser.error("Output must not overwrite a source document or canonical mapping")
    data = build_reference(args.source, args.cpdd_reference)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(f"Imported {len(data['barangays'])} barangays and {data['city_summary']['households_2020']:,} households to {args.output}")


if __name__ == "__main__":
    main()
